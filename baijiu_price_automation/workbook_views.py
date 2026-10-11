"""Keep native Excel filters and scroll positions aligned with actual quotes."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from update_baijiu_prices import build_brand_column_map, build_summary_column_map
from workbook_manual import sheet_paths, q


def refresh_views(path):
    path = Path(path)
    book = load_workbook(path, data_only=False)
    changes = {}
    with ZipFile(path) as original:
        paths = sheet_paths(original)
        patched = {}
        for name in ('茅台', '五粮液', '泸州老窖', '习酒', 'Sheet1'):
            if name not in book: continue
            sheet = book[name]
            columns = set((build_summary_column_map(sheet) if name == 'Sheet1' else build_brand_column_map(sheet)).values())
            header = 3 if name == 'Sheet1' else 2
            end = max((row for row in range(header + 1, sheet.max_row + 1)
                       if any(isinstance(sheet.cell(row, col).value, (int, float)) for col in columns)), default=header + 1)
            root = ET.fromstring(original.read(paths[name]))
            lastcol = max((c.column for c in sheet[header] if c.value is not None), default=max(columns, default=2))
            filt = root.find(q('autoFilter'))
            if filt is not None:
                # Preserve column IDs and criteria by keeping the existing start column.
                start = filt.get('ref', 'A2').split(':')[0]
                filt.set('ref', f'{start}:{get_column_letter(lastcol)}{end}')
            view = root.find(q('sheetViews') + '/' + q('sheetView'))
            top = f'A{max(header + 1, end - 6)}'
            if view is not None:
                pane = view.find(q('pane'))
                firstcol = 'B' if pane is not None and float(pane.get('xSplit', 0)) else 'A'
                top = f'{firstcol}{max(header + 1, end - 6)}'
                if pane is not None: pane.set('topLeftCell', top)
                else: view.set('topLeftCell', top)
                active = f'{get_column_letter(min(columns, default=2))}{end}'
                selections = view.findall(q('selection'))
                target = next((s for s in selections if s.get('pane') == (pane.get('activePane') if pane is not None else None)), selections[-1] if selections else None)
                if target is not None: target.set('activeCell', active); target.set('sqref', active)
                view.set('zoomScale', '85')
            patched[paths[name]] = ET.tostring(root, encoding='utf-8')
            changes[name] = {'last_quote_row': end, 'filter': filt.get('ref') if filt is not None else None, 'top_left': top}
        temp = path.with_name(path.name + '.views.tmp')
        with ZipFile(temp, 'w', ZIP_DEFLATED) as out:
            for info in original.infolist(): out.writestr(info, patched.get(info.filename, original.read(info.filename)))
    book.close()
    temp.replace(path)
    return changes
