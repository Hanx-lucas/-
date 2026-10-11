import argparse
import csv
import json
import re
import shutil
import sys
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from html import unescape
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter


DATE_RE = re.compile(r"(?:(20\d{2})[年/\-.])?\s*(\d{1,2})\s*[月/\-.]\s*(\d{1,2})\s*日?")
PRICE_RE = re.compile(r"(?<!\d)(-?\d{2,5})(?!\d)")


@dataclass
class Article:
    path: Path
    account: str
    trade_date: date
    title: str
    text: str


def norm_text(value) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return (
        text.replace("\u00a0", " ")
        .replace("（", "(")
        .replace("）", ")")
        .replace("°", "度")
        .strip()
    )


def compact(value: str) -> str:
    return re.sub(r"\s+", "", norm_text(value)).lower()


def strip_html(text: str) -> str:
    text = unescape(text)
    text = re.sub(r"(?i)<\s*br\s*/?\s*>", "\n", text)
    text = re.sub(r"(?i)</\s*(p|div|tr|li|h\d|table)\s*>", "\n", text)
    text = re.sub(r"(?i)</\s*(td|th)\s*>", "\t", text)
    text = re.sub(r"<[^>]+>", "", text)
    return text


def read_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            raw = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raw = data.decode("utf-8", errors="ignore")
    if path.suffix.lower() in {".html", ".htm"} or "<html" in raw[:500].lower():
        raw = strip_html(raw)
    return "\n".join(line.strip() for line in raw.splitlines() if line.strip())


def parse_date(text: str, default_year: int) -> date | None:
    match = DATE_RE.search(text)
    if not match:
        return None
    year = int(match.group(1) or default_year)
    month = int(match.group(2))
    day = int(match.group(3))
    try:
        return date(year, month, day)
    except ValueError:
        return None


def infer_account(path: Path, text: str, accounts: list[str]) -> str | None:
    haystack = compact(" ".join(path.parts[-3:]) + "\n" + text[:2000])
    for account in accounts:
        if compact(account) in haystack:
            return account
    return None


def load_articles(articles_dir: Path, config: dict, start: date, end: date) -> list[Article]:
    articles = []
    files = [
        p
        for p in articles_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in {".txt", ".md", ".html", ".htm", ".csv"}
    ]
    for path in sorted(files):
        if path.name.lower() == "readme.md":
            continue
        text = read_text(path)
        account = infer_account(path, text, config["accounts"])
        if not account:
            print(f"[跳过] 无法识别公众号: {path}", file=sys.stderr)
            continue
        trade_date = parse_date(path.stem, end.year) or parse_date(text[:1500], end.year)
        if not trade_date:
            print(f"[跳过] 无法识别日期: {path}", file=sys.stderr)
            continue
        if not (start <= trade_date <= end):
            continue
        title = text.splitlines()[0] if text.splitlines() else path.stem
        articles.append(Article(path, account, trade_date, title, text))
    return articles


def plausible_prices(line: str) -> list[int]:
    prices = []
    for raw in PRICE_RE.findall(line):
        n = int(raw)
        if abs(n) < 100:
            continue
        if re.search(rf"{n}\s*年", line):
            continue
        if n in {500, 100, 1000} and re.search(rf"{n}\s*(ml|ML|毫升)", line):
            continue
        if 30 <= n <= 60 and "度" in line:
            continue
        if 100 <= n <= 5000:
            prices.append(n)
    return prices


def choose_price(line: str, prices: list[int]) -> int | None:
    if not prices:
        return None
    if len(prices) == 1:
        return prices[0]
    line_c = compact(line)
    if "昨日" in line_c and ("今日" in line_c or "今天" in line_c):
        return prices[1]
    if "涨跌" in line_c and len(prices) >= 2:
        return prices[1]
    if re.search(r"\d{1,2}日", line) and len(prices) >= 2:
        return prices[1]
    return prices[1]


def lines_with_context(text: str) -> list[str]:
    raw_lines = [line.strip() for line in text.splitlines() if line.strip()]
    joined = []
    for i, line in enumerate(raw_lines):
        joined.append(line)
        if i + 1 < len(raw_lines):
            joined.append(line + "\t" + raw_lines[i + 1])
        if i + 3 < len(raw_lines):
            joined.append("\t".join(raw_lines[i : i + 4]))
    return joined


def extract_prices(article: Article, config: dict) -> list[dict]:
    rows = []
    lines = [norm_text(x) for x in article.text.splitlines() if x.strip()]
    all_aliases = [compact(a) for spec in config["products"].values()
                   for a in spec.get("aliases", {}).get(article.account, []) if a != "-"]
    for product, spec in config["products"].items():
        best = None
        for alias in spec.get("aliases", {}).get(article.account, []):
            key = compact(alias)
            if alias == "-":
                continue
            for i, line in enumerate(lines):
                text = compact(line)
                if key not in text:
                    continue
                if any(term in text and term not in key for term in ("雅致版", "品鉴装", "礼盒", "12瓶")):
                    continue
                # Remove product codes, vintage, capacity and degree from price candidates.
                match = re.search("".join(re.escape(ch) + r"\s*" for ch in key), line.lower())
                if not match:
                    continue
                tail = line[match.end():]
                if not plausible_prices(tail):
                    for following in lines[i+1:i+4]:
                        if any(a in compact(following) for a in all_aliases):
                            break
                        tail += "\t" + following
                price = choose_price(tail, plausible_prices(tail))
                if price is not None:
                    best = {"date": article.trade_date.isoformat(), "account": article.account,
                            "product": product, "alias": alias, "price": price,
                            "source_file": str(article.path), "evidence": line[:180]}
                    break
            if best:
                break
        if best:
            rows.append(best)
    return rows


def merged_value(ws, row: int, col: int):
    value = ws.cell(row, col).value
    if value is not None:
        return value
    coord = f"{get_column_letter(col)}{row}"
    for merged in ws.merged_cells.ranges:
        if coord in merged:
            return ws.cell(merged.min_row, merged.min_col).value
    return None


def build_brand_column_map(ws):
    result = {}
    current_product = None
    for col in range(1, ws.max_column + 1):
        product = merged_value(ws, 1, col)
        if product:
            current_product = norm_text(product)
        platform = norm_text(ws.cell(2, col).value)
        if current_product and platform and platform not in {"日期", "单日各平台酒价均值"}:
            result[(current_product, platform)] = col
    return result


def build_summary_column_map(ws):
    result = {}
    for col in range(1, ws.max_column + 1):
        product = norm_text(ws.cell(3, col).value or ws.cell(1, col).value)
        platform = norm_text(merged_value(ws, 2, col))
        if product and platform and platform not in {"日期", "600519.SH"}:
            result[(product, platform)] = col
    return result


def find_start_date(ws, date_col: int, start_row: int) -> tuple[date, int]:
    for row in range(start_row, ws.max_row + 1):
        value = ws.cell(row, date_col).value
        if isinstance(value, datetime):
            return value.date(), row
        if isinstance(value, date):
            return value, row
    raise ValueError(f"{ws.title}: 找不到起始日期")


def data_end_row(ws, date_col: int, start_row: int) -> int:
    marker = None
    for row in range(start_row, ws.max_row + 1):
        values = [norm_text(ws.cell(row, col).value) for col in range(1, min(ws.max_column, 10) + 1)]
        if any("各渠道酒价名称对照汇总表" in v for v in values):
            marker = row
            break
    if marker:
        return marker - 1
    last = start_row
    for row in range(start_row, ws.max_row + 1):
        if any(ws.cell(row, col).value is not None for col in range(1, ws.max_column + 1)):
            last = row
    return last


def copy_row(ws, src_row: int, dst_row: int):
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for col in range(1, ws.max_column + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        value = src.value
        if isinstance(value, str) and value.startswith("="):
            dst.value = Translator(value, origin=src.coordinate).translate_formula(dst.coordinate)
        else:
            dst.value = None


def ensure_date_row(ws, target_date: date, date_col: int, start_row: int) -> int:
    start_date, first_row = find_start_date(ws, date_col, start_row)
    target_row = first_row + (target_date - start_date).days
    if target_row < first_row:
        raise ValueError(f"{ws.title}: 目标日期早于表内起始日期 {start_date}")
    end_row = data_end_row(ws, date_col, first_row)
    while target_row > end_row:
        insert_at = end_row + 1
        if insert_at <= ws.max_row:
            ws.insert_rows(insert_at)
        copy_row(ws, end_row, insert_at)
        ws.cell(insert_at, date_col).value = start_date + timedelta(days=insert_at - first_row)
        ensure_average_formulas(ws, insert_at)
        end_row += 1
    cell = ws.cell(target_row, date_col)
    if not isinstance(cell.value, str):
        cell.value = target_date
    return target_row


def ensure_average_formulas(ws, row):
    if ws.title == "Sheet1":
        mapping = build_summary_column_map(ws)
        for col in range(3, ws.max_column + 1):
            if norm_text(merged_value(ws, 2, col)) != "单日各平台酒价均值":
                continue
            product = norm_text(ws.cell(3, col).value)
            refs = [f"{get_column_letter(c)}{row}" for (p, a), c in mapping.items()
                    if p == product and a != "单日各平台酒价均值"]
            if refs:
                args = ",".join(refs)
                ws.cell(row, col).value = f'=IF(COUNT({args})=0,"",AVERAGE({args}))'
    else:
        first = 2
        for col in range(2, ws.max_column + 1):
            if norm_text(ws.cell(2, col).value) == "单日各平台酒价均值":
                refs = f"{get_column_letter(first)}{row}:{get_column_letter(col-1)}{row}"
                ws.cell(row, col).value = f'=IF(COUNT({refs})=0,"",AVERAGE({refs}))'
                first = col + 1


def write_logs(rows: list[dict], missing: list[dict], output_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    matches_path = unique_output_path(output_dir / f"matches_{ts}.csv")
    missing_path = unique_output_path(output_dir / f"missing_review_{ts}.csv")
    with matches_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["date", "account", "product", "alias", "price", "sheet", "cell", "source_file", "evidence"])
        writer.writeheader()
        writer.writerows(rows)
    with missing_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["date", "account", "product", "reason"])
        writer.writeheader()
        writer.writerows(missing)
    return matches_path, missing_path


def unique_output_path(path: Path) -> Path:
    if not path.exists():
        return path
    stamp = datetime.now().strftime("%H%M%S_%f")
    return path.with_name(f"{path.stem}_{stamp}{path.suffix}")



def ensure_schema(wb, config):
    # Correct the two documented year aliases without changing historical prices.
    if "茅台" in wb:
        ws = wb["茅台"]
        for row in ws:
            for cell in row:
                if cell.value in {"26年飞天(原)", "26年飞天(散)"}:
                    if cell.column == 5 and norm_text(ws.cell(cell.row, 3).value) in {"25茅台原箱", "25茅台散瓶"}:
                        cell.value = cell.value.replace("26年飞天", "25年飞天")
    # The supplied template lacks 今日酒价 for 38度 and four products in Sheet1.
    if "泸州老窖" in wb:
        ws = wb["泸州老窖"]
        mapping = build_brand_column_map(ws)
        key = ("泸州老窖1573 38度", "今日酒价")
        if key not in mapping:
            ws.unmerge_cells("F1:G1")
            ws.insert_cols(7)
            ws.merge_cells("F1:H1")
            ws.cell(2, 7).value = "今日酒价"
            for row in range(1, ws.max_row + 1):
                ws.cell(row, 7)._style = copy(ws.cell(row, 6)._style)
                if row > 2 and isinstance(ws.cell(row, 8).value, str) and ws.cell(row, 8).value.startswith("="):
                    refs = f"F{row}:G{row}"
                    ws.cell(row, 8).value = f'=IF(COUNT({refs})=0,"",AVERAGE({refs}))'
            ws.column_dimensions["G"].width = ws.column_dimensions["F"].width
            ws.column_dimensions["H"].width = ws.column_dimensions["E"].width
    if "Sheet1" in wb:
        ws = wb["Sheet1"]
        mapping = build_summary_column_map(ws)
        # Retain all existing columns and append only missing product/channel fields.
        groups = []
        for account in config["accounts"]:
            products = [p for p, spec in config["products"].items()
                        if account in spec["aliases"] and (p, account) not in mapping]
            if products:
                groups.append((account, products))
        products = [p for p in config["products"] if (p, "单日各平台酒价均值") not in mapping]
        if products:
            groups.append(("单日各平台酒价均值", products))
        for account, products in groups:
            start = ws.max_column + 1
            for product in products:
                col = ws.max_column + 1
                for row in range(1, ws.max_row + 1):
                    ws.cell(row, col)._style = copy(ws.cell(row, 49 if account == "单日各平台酒价均值" else 3)._style)
                ws.cell(1, col).value = product
                ws.cell(3, col).value = product
                ws.column_dimensions[get_column_letter(col)].width = max(ws.column_dimensions["C"].width or 16, len(product) * 2 + 2)
            ws.cell(2, start).value = account
            if ws.max_column > start:
                ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=ws.max_column)


def update_workbook(workbook_path: Path, output_path: Path, rows: list[dict], config: dict, dry_run: bool):
    if dry_run:
        wb = load_workbook(workbook_path)
    else:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(workbook_path, output_path)
        wb = load_workbook(output_path)
    ensure_schema(wb, config)
    brand_maps = {name: build_brand_column_map(wb[name]) for name in ["茅台", "五粮液", "泸州老窖", "习酒"] if name in wb.sheetnames}
    summary_map = build_summary_column_map(wb["Sheet1"]) if "Sheet1" in wb.sheetnames else {}
    written = []
    row_cache = {}
    for row in rows:
        product = row["product"]
        account = row["account"]
        price = row["price"]
        trade_date = date.fromisoformat(row["date"])
        sheet_name = config["products"][product]["sheet"]
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            cache_key = (sheet_name, trade_date)
            target_row = row_cache.get(cache_key)
            if target_row is None:
                target_row = ensure_date_row(ws, trade_date, 1, 3)
                row_cache[cache_key] = target_row
            col = brand_maps.get(sheet_name, {}).get((product, account))
            if col:
                if not dry_run:
                    ws.cell(target_row, col).value = price
                    ensure_average_formulas(ws, target_row)
                item = dict(row)
                item.update({"sheet": sheet_name, "cell": ws.cell(target_row, col).coordinate})
                written.append(item)
        if "Sheet1" in wb.sheetnames:
            ws = wb["Sheet1"]
            cache_key = ("Sheet1", trade_date)
            target_row = row_cache.get(cache_key)
            if target_row is None:
                target_row = ensure_date_row(ws, trade_date, 2, 4)
                row_cache[cache_key] = target_row
            col = summary_map.get((product, account))
            if col:
                if not dry_run:
                    ws.cell(target_row, col).value = price
                    ensure_average_formulas(ws, target_row)
                item = dict(row)
                item.update({"sheet": "Sheet1", "cell": ws.cell(target_row, col).coordinate})
                written.append(item)
    if not dry_run:
        wb.save(output_path)
        from workbook_manual import normalize_manual
        normalize_manual(output_path, output_path, config)
        from workbook_views import refresh_views
        refresh_views(output_path)
    return written


def expected_missing(config: dict, start: date, end: date, extracted: list[dict]) -> list[dict]:
    seen = {(r["date"], r["account"], r["product"]) for r in extracted}
    missing = []
    current = start
    while current <= end:
        for product, spec in config["products"].items():
            for account in spec.get("aliases", {}):
                if (current.isoformat(), account, product) not in seen:
                    missing.append({"date": current.isoformat(), "account": account, "product": product, "reason": "文章未提供或未匹配到价格"})
        current += timedelta(days=1)
    return missing


def main():
    parser = argparse.ArgumentParser(description="按需求文档自动解析近七天白酒公众号价格并更新 Excel 数据库。")
    parser.add_argument("--workbook", default="../名酒酒价各平台数据库20260516.xlsx", help="原始 Excel 数据库路径")
    parser.add_argument("--articles-dir", default="article_inputs", help="文章文本/HTML目录")
    parser.add_argument("--config", default="config/baijiu_price_config.json", help="配置文件路径")
    parser.add_argument("--output", default="", help="输出 Excel 路径；为空则写入 outputs/ 自动命名")
    parser.add_argument("--start", default="", help="开始日期 YYYY-MM-DD；默认最近七天")
    parser.add_argument("--end", default="", help="结束日期 YYYY-MM-DD；默认今天")
    parser.add_argument("--dry-run", action="store_true", help="只解析和生成日志，不写入工作簿")
    args = parser.parse_args()

    base = Path(__file__).resolve().parent

    def resolve_input(raw: str, prefer_base: bool = False) -> Path:
        path = Path(raw)
        if path.is_absolute():
            return path
        cwd_path = (Path.cwd() / path).resolve()
        base_path = (base / path).resolve()
        if prefer_base:
            return base_path if base_path.exists() else cwd_path
        return cwd_path if cwd_path.exists() else base_path

    config_path = resolve_input(args.config, prefer_base=True)
    workbook_path = resolve_input(args.workbook)
    articles_dir = resolve_input(args.articles_dir)

    config = json.loads(config_path.read_text(encoding="utf-8"))
    end = date.fromisoformat(args.end) if args.end else date.today()
    start = date.fromisoformat(args.start) if args.start else end - timedelta(days=6)
    if start > end:
        parser.error("开始日期不能晚于结束日期")
    if not articles_dir.is_dir():
        parser.error(f"文章目录不存在: {articles_dir}")
    output_dir = base / "outputs"
    output_path = Path(args.output) if args.output else output_dir / f"{workbook_path.stem}_自动更新_{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
    if not output_path.is_absolute():
        output_path = (base / output_path).resolve()
    if not args.dry_run:
        output_path = unique_output_path(output_path)

    articles = load_articles(articles_dir, config, start, end)
    extracted = []
    for article in articles:
        extracted.extend(extract_prices(article, config))
    groups = {}
    for row in extracted:
        groups.setdefault((row["date"], row["account"], row["product"]), []).append(row)
    extracted = [items[0] for items in groups.values() if len({r["price"] for r in items}) == 1]
    missing = expected_missing(config, start, end, extracted)
    for key, items in groups.items():
        if len({r["price"] for r in items}) > 1:
            missing.append(dict(zip(["date", "account", "product"], key), reason="同日同渠道价格冲突，未写入"))
    if not extracted:
        _, missing_path = write_logs([], missing, output_dir)
        print(f"文章数量: {len(articles)}\n匹配价格: 0\n未生成 Excel：没有可核验的价格。\n缺失复核: {missing_path}")
        return 2
    written = update_workbook(workbook_path, output_path, extracted, config, args.dry_run)
    mapped = {(r["date"], r["account"], r["product"], r["sheet"]) for r in written}
    for row in extracted:
        sheet = config["products"][row["product"]]["sheet"]
        if (row["date"], row["account"], row["product"], sheet) not in mapped:
            missing.append({"date": row["date"], "account": row["account"], "product": row["product"],
                            "reason": f"模板 {sheet} 无此公众号列，未写入"})
    matches_path, missing_path = write_logs(written, missing, output_dir)
    if not args.dry_run and written and output_path.parent.resolve() == output_dir.resolve():
        from database_versions import DatabaseVersions
        DatabaseVersions(output_dir).publish(output_path, source=workbook_path,
                                            start=str(start), end=str(end))

    print(f"文章数量: {len(articles)}")
    print(f"匹配价格: {len(extracted)}")
    print(f"写入单元格: {len(written)}")
    print(f"输出文件: {output_path if not args.dry_run else '(dry-run 未写入)'}")
    print(f"匹配日志: {matches_path}")
    print(f"缺失复核: {missing_path}")


if __name__ == "__main__":
    sys.exit(main())




