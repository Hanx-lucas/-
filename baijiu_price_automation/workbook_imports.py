"""Validated, independent user database inputs; original uploads are never overwritten."""
from __future__ import annotations
import base64
import hashlib
from io import BytesIO
import json
from datetime import date, datetime, timedelta
from pathlib import Path
import re
from uuid import uuid4
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from database_versions import file_lock, ZONE
from update_baijiu_prices import build_brand_column_map, build_summary_column_map, data_end_row

LIMIT = 50*1024*1024

def validate_database(raw: bytes, config: dict):
    if not raw or len(raw)>LIMIT: raise ValueError('数据库须为不超过 50MB 的 XLSX 文件')
    try:
        with ZipFile(BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>250*1024*1024: raise ValueError('数据库解压后超过 250MB')
            if 'xl/workbook.xml' not in z.namelist(): raise ValueError('文件不是有效的 Excel 数据库')
        wb=load_workbook(BytesIO(raw),data_only=False)
        cached=load_workbook(BytesIO(raw),data_only=True)
    except (BadZipFile,KeyError,TypeError,OSError) as e: raise ValueError('无法读取 XLSX，请确认文件未损坏或加密') from e
    required={'更新手册','茅台','五粮液','泸州老窖','习酒','Sheet1'}
    missing=required-set(wb.sheetnames)
    if missing: raise ValueError('数据库模板缺少工作表：'+'、'.join(sorted(missing)))
    if any(s.max_row>30000 or s.max_column>300 for s in wb): raise ValueError('数据库超出支持的模板范围（3 万行、300 列）')
    recognized=0;latest=None;channels=0
    for sheet in ['茅台','五粮液','泸州老窖','习酒']:
        ws=wb[sheet];mapping=build_brand_column_map(ws)
        products=[p for p,spec in config['products'].items() if spec['sheet']==sheet]
        absent=[p for p in products if not any(key[0]==p for key in mapping)]
        if absent: raise ValueError(sheet+' 的模板品类无法识别：'+'、'.join(absent)+'。请使用相同品类名称的数据库')
        recognized+=len(products);channels+=sum((p,a) in mapping for p in products for a in config['products'][p]['aliases'])
        date_cache={}
        def day_at(row):
            if row in date_cache:return date_cache[row]
            v=cached[sheet].cell(row,1).value or ws.cell(row,1).value
            if isinstance(v,datetime):v=v.date()
            elif isinstance(v,str):
                match=re.fullmatch(r'=\s*A(\d+)\s*\+\s*(\d+)\s*',v)
                if match and int(match[1])<row:
                    previous=day_at(int(match[1]));v=previous+timedelta(days=int(match[2])) if previous else None
                else:
                    try:v=date.fromisoformat(v[:10])
                    except ValueError:v=None
            if not isinstance(v,date):v=None
            date_cache[row]=v;return v
        for row in range(3,data_end_row(ws,1,3)+1):
            # Mapping text and empty future date scaffolds are not historical quotes.
            if not any(isinstance(ws.cell(row,col).value,(int,float)) and not isinstance(ws.cell(row,col).value,bool) for col in set(mapping.values())):continue
            d=day_at(row)
            if d is not None:latest=max(latest,d) if latest else d
    if not build_summary_column_map(wb['Sheet1']):raise ValueError('Sheet1 的汇总表头无法识别')
    if channels<40:raise ValueError('公众号列与模板差异过大，请使用相同格式的数据库')
    return {'latest_price_date':latest.isoformat() if latest else None,'products':recognized,'mapped_channel_pairs':channels,'sheets':wb.sheetnames,'sha256':hashlib.sha256(raw).hexdigest()}

class DatabaseImports:
    def __init__(self,base):
        self.base=Path(base).resolve();self.folder=self.base/'database_inputs';self.manifest=self.folder/'imports.json'
    def read(self):
        if not self.manifest.is_file():return {'records':[],'active':None}
        return json.loads(self.manifest.read_text(encoding='utf-8'))
    def options(self):
        data=self.read();items=[]
        for r in reversed(data['records']):
            p=(self.folder/r['stored_name']).resolve()
            if p.parent==self.folder.resolve() and p.is_file():
                items.append({'value':str(p.relative_to(self.base.parent)).replace('\\','/'),'label':'已导入 · '+r['original_name']+' · '+(r['latest_price_date'] or '未识别报价日期'),'metadata':r})
        return items,data.get('active')
    def clear_active(self):
        with file_lock(self.folder/'.imports.lock'):
            data=self.read();data['active']=None;self._save(data)
    def _save(self,data):
        temp=self.manifest.with_suffix('.json.tmp');temp.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8');temp.replace(self.manifest)
    def import_file(self,payload,config):
        name=str(payload.get('name',''))
        if not name or Path(name).name!=name or any(c in name for c in '<>:"/\\|?*') or Path(name).suffix.lower()!='.xlsx' or name.startswith('~$'):raise ValueError('请选择单个有效的 XLSX 数据库文件')
        if '样例' in name:raise ValueError('样例数据库不能用于正式更新')
        try:raw=base64.b64decode(payload.get('data',''),validate=True)
        except (ValueError,TypeError) as e:raise ValueError('上传内容无法解码') from e
        metadata=validate_database(raw,config)
        self.folder.mkdir(parents=True,exist_ok=True)
        stem=re.sub(r'[^\w\u4e00-\u9fff.-]','_',Path(name).stem)[:70]
        stored=datetime.now(ZONE).strftime('%Y%m%d_%H%M%S_')+uuid4().hex[:8]+'_'+stem+'.xlsx'
        dest=(self.folder/stored).resolve()
        if dest.parent!=self.folder.resolve():raise ValueError('数据库路径无效')
        with file_lock(self.folder/'.imports.lock'):
            data=self.read();record={**metadata,'stored_name':stored,'original_name':name,'imported_at':datetime.now(ZONE).isoformat()}
            dest.write_bytes(raw)
            active=str(dest.relative_to(self.base.parent)).replace('\\','/')
            data['records'].append(record);data['active']=active
            try:self._save(data)
            except Exception:dest.unlink();raise
        return {'value':active,**record}
